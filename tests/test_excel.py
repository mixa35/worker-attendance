"""Workbook rendering — the artefact the owner actually receives."""

from __future__ import annotations

import sqlite3
from pathlib import Path

import pytest
from openpyxl import load_workbook

from worker_attendance import excel as excel_module
from worker_attendance.excel import COL_ID, COL_TEAM, COL_WORKER, FIRST_DAY_COL, PRESENT_MARK


def _load(path: Path):
    return load_workbook(path)


@pytest.mark.parametrize(
    ("year", "month", "days"),
    [(2026, 2, 28), (2024, 2, 29), (2026, 4, 30), (2026, 1, 31)],
)
def test_day_columns_match_the_length_of_the_month(
    seeded: sqlite3.Connection, data_dir: Path, year: int, month: int, days: int
) -> None:
    excel_module.render_month_sheet(year, month)
    ws = _load(excel_module.XLSX_PATH)[f"{year:04d}-{month:02d}"]

    assert ws.cell(row=1, column=FIRST_DAY_COL).value == 1
    assert ws.cell(row=1, column=FIRST_DAY_COL - 1 + days).value == days
    assert ws.cell(row=1, column=FIRST_DAY_COL + days).value == "Total"


def test_sheet_is_named_year_dash_month(
    seeded: sqlite3.Connection, data_dir: Path
) -> None:
    excel_module.render_month_sheet(2026, 5)
    assert _load(excel_module.XLSX_PATH).sheetnames == ["2026-05"]


def test_present_days_are_marked_and_totalled(mark, data_dir: Path) -> None:
    mark(1, "2026-05-04")
    mark(1, "2026-05-05")
    mark(1, "2026-05-06", present=0)

    excel_module.render_month_sheet(2026, 5)
    ws = _load(excel_module.XLSX_PATH)["2026-05"]

    ana = _row_for(ws, "Ana")
    assert ws.cell(row=ana, column=FIRST_DAY_COL + 3).value == PRESENT_MARK  # 4th
    assert ws.cell(row=ana, column=FIRST_DAY_COL + 4).value == PRESENT_MARK  # 5th
    assert ws.cell(row=ana, column=FIRST_DAY_COL + 5).value is None  # 6th, absent
    assert ws.cell(row=ana, column=FIRST_DAY_COL + 31).value == 2  # Total, May = 31d


def test_attendance_in_another_month_is_not_counted(mark, data_dir: Path) -> None:
    mark(1, "2026-04-30")
    mark(1, "2026-06-01")

    excel_module.render_month_sheet(2026, 5)
    ws = _load(excel_module.XLSX_PATH)["2026-05"]
    assert ws.cell(row=_row_for(ws, "Ana"), column=FIRST_DAY_COL + 31).value == 0


def test_rendering_twice_produces_an_identical_sheet(mark, data_dir: Path) -> None:
    """render_month_sheet documents itself as idempotent — hold it to that."""
    mark(1, "2026-05-04")

    excel_module.render_month_sheet(2026, 5)
    first = _snapshot(_load(excel_module.XLSX_PATH)["2026-05"])

    excel_module.render_month_sheet(2026, 5)
    wb = _load(excel_module.XLSX_PATH)

    assert wb.sheetnames == ["2026-05"], "re-rendering must replace, not duplicate"
    assert _snapshot(wb["2026-05"]) == first


def test_each_month_gets_its_own_sheet(
    seeded: sqlite3.Connection, data_dir: Path
) -> None:
    excel_module.render_month_sheet(2026, 4)
    excel_module.render_month_sheet(2026, 5)
    assert set(_load(excel_module.XLSX_PATH).sheetnames) == {"2026-04", "2026-05"}


def test_national_id_falls_back_to_a_hashed_row_id(
    seeded: sqlite3.Connection, data_dir: Path
) -> None:
    excel_module.render_month_sheet(2026, 5)
    ws = _load(excel_module.XLSX_PATH)["2026-05"]

    assert ws.cell(row=_row_for(ws, "Ana"), column=COL_ID).value == "01001"
    assert ws.cell(row=_row_for(ws, "Beka"), column=COL_ID).value == "#2"


def test_inactive_worker_appears_only_when_the_month_has_their_attendance(
    mark, data_dir: Path
) -> None:
    """Giorgi is inactive: dropped from a clean month, kept where he has history."""
    excel_module.render_month_sheet(2026, 5)
    assert _row_for(_load(excel_module.XLSX_PATH)["2026-05"], "Giorgi") is None

    mark(3, "2026-05-04")
    excel_module.render_month_sheet(2026, 5)
    ws = _load(excel_module.XLSX_PATH)["2026-05"]

    giorgi = _row_for(ws, "Giorgi")
    assert giorgi is not None
    assert ws.cell(row=giorgi, column=COL_TEAM).value == "Bravo"


def test_workbook_is_written_atomically(
    seeded: sqlite3.Connection, data_dir: Path
) -> None:
    """The .tmp file must not survive the rename."""
    excel_module.render_month_sheet(2026, 5)
    assert excel_module.XLSX_PATH.exists()
    assert not excel_module.XLSX_PATH.with_suffix(".xlsx.tmp").exists()
    assert list(data_dir.glob("*.tmp")) == []


def _row_for(ws, worker_name: str) -> int | None:
    for row in range(2, ws.max_row + 1):
        if ws.cell(row=row, column=COL_WORKER).value == worker_name:
            return row
    return None


def _snapshot(ws) -> list[tuple]:
    return [
        tuple(cell.value for cell in row)
        for row in ws.iter_rows(min_row=1, max_row=ws.max_row)
    ]
